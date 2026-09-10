"""Editable Word reports (English), modeled on the client's samples:
- price_summary.docx    — format follows the tender's price scheme (both sample formats)
- summary_list.docx     — Stage I / Stage II narrative conclusions
- evaluation_record.docx — detailed per-bidder record with evidence and page citations
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document as Docx

from .schemas import EvaluationResult


def _money(v: float | None) -> str:
    return f"{v:,.2f}" if v is not None else "—"


def _num(v: float | None) -> str:
    return f"{v:,.2f}".rstrip("0").rstrip(".") if v is not None else "—"


def _unit_price(v: float | None) -> str:
    """At least 2 decimals, up to 4 when quoted more precisely (e.g. 1.278)."""
    if v is None:
        return "—"
    s = f"{v:,.4f}".rstrip("0")
    if len(s.partition(".")[2]) < 2:
        return f"{v:,.2f}"
    return s


def _set_cell(cell, value: str, bold: bool) -> None:
    cell.text = value
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold


def _header_row(table, labels: list[str]) -> None:
    for cell, label in zip(table.rows[0].cells, labels):
        _set_cell(cell, label, bold=True)


def _fill_row(table, values: list[str], bold: bool = False) -> None:
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        _set_cell(cell, value, bold)


def _title(doc: Docx, result: EvaluationResult, heading: str) -> None:
    doc.add_paragraph(f"Tender Ref.: {result.rubric.tender_ref}")
    if result.rubric.subject:
        doc.add_paragraph(result.rubric.subject)
    doc.add_heading(heading, level=1)


# ---------------------------------------------------------------- price summary

def _ce_table(doc: Docx, result: EvaluationResult) -> None:
    scheme = result.rubric.price_scheme
    table = doc.add_table(rows=1, cols=6, style="Table Grid")
    _header_row(table, [
        "Name of Tenderer",
        f"Estimated Goods Price (HK$)\n(M) x {scheme.quantity:,.0f} {scheme.unit}",
        f"Optimal Dosage (\"D\")\n({scheme.unit}/tonne of dried solids)",
        "One-time Unit Price (\"M\")\n(HK$/" + scheme.unit + ")",
        "Cost-effectiveness of Tender\n(D) x (M)",
        "Ranking of Tender\n(in terms of cost-effectiveness)",
    ])
    for row in result.price_rows:
        bold = row.tenderer == result.recommended
        _fill_row(table, [
            row.tenderer,
            _money(row.estimated_goods_price),
            _num(row.dosage_rounded if row.dosage_rounded is not None else row.dosage),
            _money(row.unit_price_hkd),
            _money(row.cost_effectiveness) if row.cost_effectiveness is not None else "cannot be calculated",
            str(row.ranking) if row.ranking is not None else "not applicable",
        ], bold=bold)


def _pq_table(doc: Docx, result: EvaluationResult) -> None:
    scheme = result.rubric.price_scheme
    table = doc.add_table(rows=1, cols=3, style="Table Grid")
    _header_row(table, [
        "Name of Tenderers",
        f"Unit Price per {scheme.unit.rstrip('s')} (HK$)",
        "Estimated Goods Price (HK$)",
    ])
    for row in result.price_rows:
        bold = row.tenderer == result.recommended
        price = _money(row.estimated_goods_price)
        if row.ranking is not None:
            price += f"  [{row.ranking}]"
        _fill_row(table, [row.tenderer, _unit_price(row.unit_price_hkd), price], bold=bold)


def _price_notes(doc: Docx, result: EvaluationResult) -> None:
    doc.add_paragraph()
    doc.add_paragraph("Notes:").runs[0].bold = True
    notes = ["The recommended offer is in bold type.",
             "The estimated goods price means the one-time unit price multiplied by the "
             f"estimated quantity specified in the Price Schedule "
             f"(i.e. {result.rubric.price_scheme.quantity:,.0f} {result.rubric.price_scheme.unit})."]
    if result.rubric.price_scheme.type == "cost_effectiveness":
        notes.append("A lower figure for the cost-effectiveness of a tender denotes a more "
                     "cost-effective offer.")
    else:
        notes.append("The figures in the brackets [ ] denote the ranking of tenders in terms "
                     "of price, with ranking number 1 being the lowest priced offer.")
    for row in result.price_rows:
        if row.arithmetic_ok is False:
            notes.append(f"The TAP noted an {row.remark.split('; ')[0]} in the offer of "
                         f"{row.tenderer}.")
        if row.cost_effectiveness is None and row.remark.startswith("cannot"):
            notes.append(f"The cost-effectiveness of {row.tenderer} cannot be calculated "
                         "from the information provided in its offer.")
    for i, note in enumerate(notes, 1):
        doc.add_paragraph(f"({i}) {note}")


def _save(doc, out_path: Path) -> None:
    """Write the document with one sequential write. python-docx's zip writer seeks
    back to patch headers, which network/FUSE filesystems (Cloud Run's GCS mount)
    reject as out-of-order writes and only recover through a slow fallback."""
    buf = io.BytesIO()
    doc.save(buf)
    out_path.write_bytes(buf.getvalue())


def render_price_summary(result: EvaluationResult, out_path: Path) -> None:
    doc = Docx()
    _title(doc, result, "Summary of Cost-effectiveness"
           if result.rubric.price_scheme.type == "cost_effectiveness" else "Price Summary")
    if result.rubric.price_scheme.type == "cost_effectiveness":
        _ce_table(doc, result)
    else:
        _pq_table(doc, result)
    _price_notes(doc, result)
    _save(doc, out_path)


# ---------------------------------------------------------------- summary list

def render_summary_list(result: EvaluationResult, out_path: Path) -> None:
    doc = Docx()
    _title(doc, result, "Summary of Tender Evaluation")
    doc.add_heading("Stage I – Completeness Check", level=2)
    doc.add_paragraph(result.stage1_conclusion)
    doc.add_heading("Stage II – Assessment of Compliance with Essential Requirements", level=2)
    doc.add_paragraph(result.stage2_conclusion)
    if result.recommended:
        doc.add_heading("Recommendation", level=2)
        best = next(r for r in result.price_rows if r.tenderer == result.recommended)
        metric = ("a cost-effectiveness of HK$" + _money(best.cost_effectiveness)
                  if best.cost_effectiveness is not None
                  else "a total estimated amount of HK$" + _money(best.estimated_goods_price))
        doc.add_paragraph(
            f"Based on the above evaluation results, the TAP found that {result.recommended}'s "
            f"offer fully complied with all the procedural and essential requirements, and "
            f"recommended acceptance of {result.recommended}'s offer with {metric}.")
    _save(doc, out_path)


# ---------------------------------------------------------------- evaluation record

def render_evaluation_record(result: EvaluationResult, out_path: Path) -> None:
    doc = Docx()
    _title(doc, result, "Detailed Evaluation Record")

    doc.add_heading("Stage I – Completeness Check", level=2)
    items = result.rubric.stage1_checklist
    table = doc.add_table(rows=1, cols=2 + len(items), style="Table Grid")
    _header_row(table, ["Tenderer"] + [i.id for i in items] + ["Result"])
    for r in result.stage1:
        cells = [r.tenderer]
        for i in items:
            f = r.presence.get(i.id)
            mark = "✓" if (f and f.present) else "✗"
            if f and f.page:
                mark += f" (p.{f.page})"
            cells.append(mark)
        cells.append("Pass" if r.passed else "Fail")
        _fill_row(table, cells)
    doc.add_paragraph("Checklist items:")
    for i in items:
        src = f" — {i.source_clause}" if i.source_clause else ""
        doc.add_paragraph(f"{i.id}: {i.item}{src}", style="List Bullet")

    doc.add_heading("Stage II – Essential Requirements", level=2)
    for r in result.stage2:
        doc.add_paragraph(f"{r.tenderer} — {'Pass' if r.passed else 'Fail'}").runs[0].bold = True
        table = doc.add_table(rows=1, cols=4, style="Table Grid")
        _header_row(table, ["Requirement", "Finding", "Evidence", "Page"])
        reqs = {q.id: q.requirement for q in result.rubric.stage2_requirements}
        for req_id, f in r.findings.items():
            _fill_row(table, [
                f"{req_id}: {reqs.get(req_id, '')}",
                f.complies,
                f.evidence or "—",
                str(f.page) if f.page else "—",
            ])
        doc.add_paragraph()
    _save(doc, out_path)


def render_all(result: EvaluationResult, out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = [out_dir / "price_summary.docx", out_dir / "summary_list.docx",
             out_dir / "evaluation_record.docx"]
    render_price_summary(result, paths[0])
    render_summary_list(result, paths[1])
    render_evaluation_record(result, paths[2])
    return paths
