"""Word deliverables: generated, reopenable, and carrying the right numbers."""
from docx import Document as Docx

from app.evaluate import evaluate
from app.report import render_all


def _cells(table):
    return [[c.text for c in row.cells] for row in table.rows]


def test_reports_generated_and_correct(synthetic_case, tmp_path):
    rubric, bids = synthetic_case
    result = evaluate(rubric, bids)
    paths = render_all(result, tmp_path)
    assert [p.name for p in paths] == [
        "price_summary.docx", "summary_list.docx", "evaluation_record.docx"]
    assert all(p.is_file() and p.stat().st_size > 0 for p in paths)

    # Price summary: cost-effectiveness table with rounded dosage and CE figures.
    price = Docx(str(paths[0]))
    rows = _cells(price.tables[0])
    assert "Name of Tenderer" in rows[0][0]
    flat = {cell for row in rows[1:] for cell in row}
    assert "39.60" in flat                 # Bidder B CE
    assert "56.00" in flat                 # Bidder A CE (2.8 x 20.00, rounded dosage)
    assert "cannot be calculated" in flat  # Bidder D
    assert "not applicable" in flat        # Bidder D ranking

    # Recommended offer is in bold type (sample convention).
    b_row = next(r for r in price.tables[0].rows if r.cells[0].text == "Bidder B")
    assert all(run.bold for cell in b_row.cells for p in cell.paragraphs for run in p.runs)

    # Summary list: both stage conclusions + recommendation paragraph.
    summary_text = "\n".join(p.text for p in Docx(str(paths[1])).paragraphs)
    assert "Stage I – Completeness Check" in summary_text
    assert "Bidder B's offer" in summary_text and "recommended acceptance" in summary_text

    # Evaluation record: Stage I matrix marks Bidder C's missing certificate.
    record = Docx(str(paths[2]))
    s1_rows = _cells(record.tables[0])
    c_row = next(r for r in s1_rows if r[0] == "Bidder C")
    assert "✗" in c_row[4] and c_row[-1] == "Fail"
